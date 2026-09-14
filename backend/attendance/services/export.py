"""
Genera la bitácora mensual en formato Excel o PDF, para que se pueda
descargar e imprimir igual que la planilla que ya usa la empresa.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

MESES_ES = [
    "", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
]

HEADERS = ["Empleado", "Documento", "Área", "Tipo", "Fecha", "Hora", "Método", "Confianza"]


def _row_data(log):
    local_ts = log.timestamp
    return [
        log.employee.full_name,
        log.employee.document_id,
        log.employee.department or "-",
        log.get_log_type_display(),
        local_ts.strftime("%Y-%m-%d"),
        local_ts.strftime("%H:%M:%S"),
        log.get_method_display(),
        f"{log.match_confidence:.0%}" if log.match_confidence is not None else "-",
    ]


def build_monthly_excel(logs, year: int, month: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bitácora"

    title = f"Bitácora de ingreso y salida de personal - {MESES_ES[month]} {year}"
    ws.merge_cells(f"A1:{get_column_letter(len(HEADERS))}1")
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")

    header_row = 3
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F6E56", end_color="1F6E56", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for i, log in enumerate(logs, start=header_row + 1):
        for col, value in enumerate(_row_data(log), start=1):
            ws.cell(row=i, column=col, value=value)

    for col in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_monthly_pdf(logs, year: int, month: int) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(letter),
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(f"Bitácora de ingreso y salida de personal - {MESES_ES[month]} {year}", styles["Title"]),
        Spacer(1, 12),
    ]

    data = [HEADERS] + [_row_data(log) for log in logs]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F6E56")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F3")]),
    ]))
    elements.append(table)

    doc.build(elements)
    return buffer.getvalue()
